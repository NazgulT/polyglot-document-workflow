# app/extractors/docx.py

import io

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from app.extractors.base import TextExtractor, ExtractionError


class DocxExtractor:
    """
    Extract text from .docx bytes using python-docx.

    We walk paragraphs and tables separately so text inside table
    cells isn't silently dropped — a common bug in naive implementations.
    """

    def extract(self, data: bytes) -> str:
        try:
            doc = DocxDocument(io.BytesIO(data))
        except Exception as exc:
            raise ExtractionError(f"Failed to open DOCX: {exc}") from exc

        parts: list[str] = []

        # Body paragraphs (headings, normal text, list items)
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # Table cells — iterate rows and cells explicitly
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        parts.append(text)

        full_text = "\n".join(parts).strip()

        if not full_text:
            raise ExtractionError("DOCX produced no extractable text.")

        return full_text